"""The diamond register (``DIAMOND 31.xlsx``): a pivot, unrolled into lines.

Built to the layout ``build_dia.py`` documents, because the workbook itself is
not yet on hand: sheet ``Sheet``, columns A category · B batch · C item code ·
D unused · E size · F weight. When the real file arrives, this parse (and its
fixture) is the only thing that should need to change.
"""
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import load_workbook

SHEET = "Sheet"

CATEGORY_NAMES = {"Diamond": "Natural Diamond", "HPHT Diamond": "HPHT Lab Grown", "Lab Grown": "Lab Grown (CVD?)",
                  "Solitare": "Solitaire", "Foil Polki": "Foil Polki"}


@dataclass
class Row:
    src: str
    category: str
    batch_no: str
    item_code: str
    size_text: str
    ct: Decimal


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _sheet(workbook):
    return workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]


def header_problems(fileobj):
    """Why this is not the diamond register, or ``[]``."""
    try:
        workbook = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as error:
        return [f"It does not open as a workbook ({error})."]
    for values in _sheet(workbook).iter_rows(max_row=10, values_only=True):
        cells = [_text(v) for v in values[:2]] + [""] * 2
        if cells[0] == "Raw Material" or cells[1] == "Batch No":
            return []
    return ["No 'Raw Material' / 'Batch No' header in the first ten rows."]


def parse(fileobj):
    sheet = _sheet(load_workbook(fileobj, read_only=True, data_only=True))
    current = {"category": "", "batch_no": "", "item_code": "", "size_text": ""}
    rows = []
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = list(values) + [None] * 6
        a, b, c, d, e = (_text(v) for v in values[:5])
        weight = values[5]
        if a == "Raw Material" or b == "Batch No":
            continue
        if "Total" in a + b + c + d:
            continue
        if not (a or b or c or d or e):
            continue            # the pivot's unlabelled grand totals
        if a:
            current["category"] = CATEGORY_NAMES.get(a, a)
        if b:
            current["batch_no"] = b
        if c:
            current["item_code"] = c
        if e:
            current["size_text"] = e
        if isinstance(weight, bool) or not isinstance(weight, (int, float)):
            continue
        rows.append(Row(src=f"{sheet.title}!{number}", ct=Decimal(str(weight)).quantize(Decimal("0.01")), **current))
    return rows
