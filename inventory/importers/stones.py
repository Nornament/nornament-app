"""The stones register: ``Nornament_Stones_Stock_Chetna_*.xlsx``.

Five sheets are stones (SP, PR, Mix, SR, SL). Googri and Jadau are finished
components, not stones, and are skipped. Row 1 is the header; 14 columns
matter and the rest (``New Code``, ``Old Gati Code``) are dead. Values are kept
as typed. Cleaning them up is a data-quality job, and the review screen is
where a human decides anything this reader cannot.
"""
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import load_workbook

SHEETS = ("SP", "PR", "Mix", "SR", "SL")

COLUMNS = {
    "batch": "New Gati Code",
    "pouch_no": "Batch No.",
    "carton": "Box No",
    "category": "Category",
    "stone_name": "Stone Name",
    "shape": "Shape",
    "size_text": "Size Length * Width",
    "colour": "Colour",
    "cut": "Cut",
    "quality": "Quality",
    "pcs": "Pcs",
    "ct": "Weight in Cts / Qty",
    "rate": "Price Per Carat / Pc",
    "remarks": "Remarks",
}

REQUIRED = ("New Gati Code", "Weight in Cts / Qty")


@dataclass
class Row:
    src: str
    batch: str
    pouch_no: str
    carton: str
    category: str
    stone_name: str
    shape: str
    size_text: str
    colour: str
    cut: str
    quality: str
    remarks: str
    pcs: int = None
    ct: Decimal = None
    rate: Decimal = None


def _text(value):
    """A cell as the owner typed it. ``1.0`` from Excel is the pouch number 1."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value):
    """Numbers only: "N.A" in a weight column is no weight, not an error."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return Decimal(str(value))


def _pieces(value):
    number = _number(value)
    # ponytail: a fractional piece count is rounded; flag it in review if one ever appears
    return None if number is None else int(number.to_integral_value())


def _sheets(workbook):
    return [name for name in SHEETS if name in workbook.sheetnames]


def _header(sheet):
    return [_text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])]


def header_problems(fileobj):
    """Why this is not a stones register, or ``[]``."""
    try:
        workbook = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as error:
        return [f"It does not open as a workbook ({error})."]
    names = _sheets(workbook)
    if not names:
        return [f"None of the stone sheets ({', '.join(SHEETS)}) is in it."]
    problems = []
    for name in names:
        header = _header(workbook[name])
        missing = [column for column in REQUIRED if column not in header]
        if missing:
            problems.append(f"Sheet {name} has no {', '.join(missing)} column.")
    return problems


def parse(fileobj):
    workbook = load_workbook(fileobj, read_only=True, data_only=True)
    rows = []
    for name in _sheets(workbook):
        sheet = workbook[name]
        header = _header(sheet)
        where = {field: header.index(title) for field, title in COLUMNS.items() if title in header}
        for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(v not in (None, "") for v in values):
                continue

            def cell(field):
                index = where.get(field)
                return values[index] if index is not None and index < len(values) else None

            rows.append(Row(
                src=f"{name}!{number}",
                batch=_text(cell("batch")).upper(),
                pouch_no=_text(cell("pouch_no")),
                carton=_text(cell("carton")),
                category=_text(cell("category")),
                stone_name=_text(cell("stone_name")),
                shape=_text(cell("shape")),
                size_text=_text(cell("size_text")),
                colour=_text(cell("colour")),
                cut=_text(cell("cut")),
                quality=_text(cell("quality")),
                remarks=_text(cell("remarks")),
                pcs=_pieces(cell("pcs")),
                ct=_number(cell("ct")),
                rate=_number(cell("rate")),
            ))
    return rows
