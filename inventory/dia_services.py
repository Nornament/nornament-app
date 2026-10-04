"""Diamond reads that every screen shares, and every diamond write.

Carats on hand are the sum of a line's movements (the stones' rule). Cost is
the line's own cost when a purchase or an assortment set one, else the
inventory's own rate card — the latest rate for the line's item code and size,
else for the code at any size — never the stock app's chart. Sale is always
the rate card's.
"""
from decimal import Decimal, InvalidOperation

from django.contrib.auth.models import Group, Permission
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import OuterRef, Subquery
from django.utils import timezone
from openpyxl import load_workbook

from accounts.capabilities import INV_MASTERS, VIEW_COST, VIEW_SALE, VIEW_VENDOR
from stock.models import Vendor
from stock.services import ServiceError, log, require

from . import dia_rules, inputs
from .dia_rules import canonical_code
from .dia_seed import COLOUR_LADDER
from .models import DiamondCode, DiamondLine, DiamondLineCost, DiamondRate, DiamondTerm, Movement, StockTake, balance

#: the prototype's seven rights, as the permissions they are, plus part 2's Record stock movements
RIGHTS = [("view_cost", "See cost"), ("view_sale", "See sale"), ("view_margin", "See margin"),
          ("inv_purchase", "Post purchase"), ("inv_job", "Job cards"), ("inv_assort", "Assort"),
          ("inv_move", "Record stock movements"), ("inv_masters", "Edit settings")]
ROLE_ORDER = ["ADMIN", "ACCOUNTS", "SALES", "GRAPHIC", "PRODUCTION", "KARIGAR"]

#: IVY export: header on row 3, diamond band columns by position (the names repeat across bands)
IVY_HEADER_ROW, IVY_CODE, IVY_SIZE, IVY_COST, IVY_SALE = 3, 20, 26, 32, 33


def _by(user):
    return user if getattr(user, "is_authenticated", False) else None


def _own_costs(line):
    """A line's own cost rows in force today, latest first: a row dated ahead waits for its day."""
    return (DiamondLineCost.objects.filter(line=line, effective_from__lte=timezone.localdate())
            .order_by("-effective_from", "-pk"))


def stocked_lines(queryset=None):
    """Lines with ``on_ct`` (the sum of their movements) and ``own_cost`` (their latest own cost per
    carat, or ``None`` for a line that never had one)."""
    queryset = DiamondLine.objects.all() if queryset is None else queryset
    return queryset.select_related(
        "category", "band", "shape_override", "colour_override", "code__shape", "code__colour", "code__clarity"
    ).annotate(
        on_ct=balance("ct"), own_cost=Subquery(_own_costs(OuterRef("pk")).values("cost_rate")[:1])
    ).order_by("pk")


def rate_table():
    """The current rate per (code, size), per field: later effective dates, then later rows,
    win (a rate dated in the future waits for its day), but only for the field a row actually sets — a cost-only row never hides an earlier
    sale rate."""
    table = {}
    for rate in DiamondRate.objects.filter(effective_from__lte=timezone.localdate()).order_by("effective_from", "pk"):
        entry = table.setdefault((rate.code_id, rate.size_text), {"cost": None, "sale": None})
        if rate.cost_rate is not None:
            entry["cost"] = rate.cost_rate
        if rate.sale_rate is not None:
            entry["sale"] = rate.sale_rate
    return table


def price(line, rates):
    """(cost, sale) per carat; ``None`` is "not set".

    Cost is the line's own latest cost when it has one, else the rate card's for its code and
    size, else for its code at any size. Sale is the rate card's alone. A line read through
    ``stocked_lines`` carries ``own_cost``; any other is looked up, so no caller can skip it.
    """
    exact = rates.get((line.code_id, line.size_text), {})
    any_size = rates.get((line.code_id, ""), {})
    own = line.own_cost if hasattr(line, "own_cost") else _own_costs(line).values_list("cost_rate", flat=True).first()
    card = exact.get("cost") if exact.get("cost") is not None else any_size.get("cost")
    cost = own if own is not None else card
    sale = exact.get("sale") if exact.get("sale") is not None else any_size.get("sale")
    return cost, sale


def term(kind, value):
    found, _ = DiamondTerm.objects.get_or_create(kind=kind, value=value)
    return found


def _last_ref_number():
    # ponytail: read-the-max under the caller's transaction, as for pouches
    last = DiamondLine.objects.order_by("-ref").values_list("ref", flat=True).first()
    return int(last[4:]) if last else 0


def new_line(**fields):
    """A line with the next ``NRD-`` reference and nothing in it: its carats arrive by the movement
    its caller posts (a purchase or an assortment), in the caller's transaction."""
    inputs.fits(DiamondLine, **fields)
    return DiamondLine.objects.create(ref=f"NRD-{_last_ref_number() + 1:06d}", **fields)


def line_label(line):
    """How a line is named wherever it is picked or listed: ref · item code · size · batch."""
    return f"{line.ref} · {line.code_id} · {line.size_text or line.band.value} · {line.batch_no or 'no batch'}"


def line_choices(queryset=None):
    """The lines with carats on hand, for a picker."""
    return [{"pk": line.pk, "label": line_label(line), "ct": line.on_ct}
            for line in stocked_lines(queryset).filter(on_ct__gt=0)]


def term_values(kind):
    """A master list's values for a picker, in its order. An unresolved ``?`` or ``(`` value is
    never offered: it is a reading to correct, not a grade to choose."""
    values = DiamondTerm.objects.filter(kind=kind).order_by("sort", "value").values_list("value", flat=True)
    return [value for value in values if not value.startswith(("?", "("))]


def listed(kind, value):
    """The master-list term for a typed value. A purchase or an assortment never adds to the lists:
    Settings does."""
    value = (value or "").strip()
    if not value:
        return None
    found = DiamondTerm.objects.filter(kind=kind, value=value).first()
    if found is None:
        raise ServiceError(f"{value} is not on the {kind} list; add it in Settings first.")
    return found


def _code_name(shape, colour, clarity):
    """A new code's name in the register's own grammar (``DRFGH VS-SI``) when the shape and colour
    have tokens, so it reads back the same; otherwise the plain words (``D Polki I-J SI-I``)."""
    shapes = {name: token for token, name in dia_rules.SHAPES}                 # OV after OVL: OV wins
    colours = {name: token for token, name in reversed(dia_rules.COLOURS)}     # KL and MN, never LC or LB
    if shape in shapes and (colour in colours or colour in COLOUR_LADDER):
        return canonical_code(f"D{shapes[shape]}{colours.get(colour, colour)} {clarity}")
    return canonical_code(f"D {shape} {colour} {clarity}")


def code_for(user, shape, colour, clarity):
    """The item code that means exactly this shape, colour and clarity — an existing one, a
    confirmed one first — else a new one, created confirmed (owner, 2026-10-01).

    Deviation from the brief (controller ruling, 2026-10-01): clarity may be blank only when
    colour is a Fancy colour (a colour term whose value starts with "Fancy") — it then finds or
    makes the code for that shape and colour with no clarity. A blank clarity with any other
    colour is refused, same as a blank shape or colour.
    """
    shape, colour, clarity = ((value or "").strip() for value in (shape, colour, clarity))
    if not shape or not colour or (not clarity and not colour.startswith("Fancy")):
        raise ServiceError("Give the shape, colour and clarity.")
    shape = listed(DiamondTerm.SHAPE, shape)
    colour = listed(DiamondTerm.COLOUR, colour)
    clarity = listed(DiamondTerm.CLARITY, clarity) if clarity else None
    found = (DiamondCode.objects.filter(shape=shape, colour=colour, clarity=clarity)
             .order_by("-confirmed", "item_code").first())
    if found is not None:
        return found
    name = _code_name(shape.value, colour.value, clarity.value if clarity else "")
    inputs.fits(DiamondCode, item_code=name)
    if DiamondCode.objects.filter(pk=name).exists():
        raise ServiceError(f"{name} already reads differently; correct it in Settings first.")
    code = DiamondCode.objects.create(item_code=name, shape=shape, colour=colour, clarity=clarity, confirmed=True)
    log(user, "INSERT", "inv_dia_code", code.pk, f"{code.pk} from a description")
    return code


def sized(size_text, pcs=None, ct=None):
    """A size's band (part 3's rules). A size that fits no band is a carat band at its weight per
    stone, when there are pieces to divide by (the owner, 2026-09-30)."""
    out = dia_rules.size_band(size_text)
    if out.band == "?" and (pcs or 0) > 0 and (ct or 0) > 0:
        per_stone = (ct / pcs).quantize(Decimal("0.001"))
        return dia_rules.Sized("carat band", per_stone, per_stone)
    return out


@transaction.atomic
def open_lines(user, specs, import_batch=None):
    """New lines, each with an Opening Balance of its carats."""
    require(user, INV_MASTERS, "Only a role that edits inventory records can add stock.")
    number, now, by = _last_ref_number(), timezone.now(), _by(user)
    lines, weights, pieces = [], [], []
    for spec in specs:
        spec = dict(spec)
        ct = spec.pop("ct")
        pcs = spec.pop("pcs", None)
        if ct is not None and ct < 0:
            raise ServiceError("A weight cannot be negative.")
        number += 1
        lines.append(DiamondLine(ref=f"NRD-{number:06d}", import_batch=import_batch, **spec))
        weights.append(ct)
        pieces.append(pcs)
    DiamondLine.objects.bulk_create(lines)
    Movement.objects.bulk_create([
        Movement(diamond=line, reason=Movement.Reason.OPENING_BALANCE, direction=Movement.IN,
                 pcs=pcs, ct=ct, occurred_at=now, recorded_by=by, ref=line.src)
        for line, ct, pcs in zip(lines, weights, pieces)
    ])
    if lines:
        log(user, "INSERT", "inv_dia_line", f"{lines[0].ref}..{lines[-1].ref}", f"{len(lines)} diamond lines opened")
    return lines


def recount_line(user, line, ct, note=""):
    require(user, INV_MASTERS, "Only a role that edits inventory records can recount.")
    if ct is None or ct < 0:
        raise ServiceError("A counted weight has to be zero or more.")
    held = stocked_lines(DiamondLine.objects.filter(pk=line.pk)).get().on_ct or Decimal("0")
    if ct == held:
        return None
    move = Movement.objects.create(
        diamond=line, reason=Movement.Reason.RECOUNT_ADJUSTMENT, direction=Movement.IN if ct > held else Movement.OUT,
        ct=abs(ct - held), note=note, recorded_by=_by(user),
    )
    log(user, "INSERT", "inv_movement", line.pk, f"recount of {line.ref}: {held} → {ct} ct")
    return move


def _in_use(t):
    ranges = DiamondTerm.objects.filter(kind=t.kind).exclude(pk=t.pk).exclude(expands_to="")
    if any(t.value in other.expands_to.split() for other in ranges):
        return True
    if t.kind == DiamondTerm.CATEGORY:
        return DiamondLine.objects.filter(category=t).exists() or StockTake.objects.filter(category=t).exists()
    if t.kind == DiamondTerm.BAND:
        return DiamondLine.objects.filter(band=t).exists()
    if t.kind == DiamondTerm.SHAPE:
        return DiamondCode.objects.filter(shape=t).exists() or DiamondLine.objects.filter(shape_override=t).exists()
    if t.kind == DiamondTerm.COLOUR:
        return DiamondCode.objects.filter(colour=t).exists() or DiamondLine.objects.filter(colour_override=t).exists()
    return DiamondCode.objects.filter(clarity=t).exists()


def add_term(user, kind, value):
    require(user, INV_MASTERS, "Only a role that edits settings can add a value.")
    value = (value or "").strip()
    if kind not in dict(DiamondTerm.KINDS) or not value:
        raise ServiceError("A value needs a list and a name.")
    if DiamondTerm.objects.filter(kind=kind, value=value).exists():
        raise ServiceError(f"{value} is already on the list.")
    created = DiamondTerm.objects.create(kind=kind, value=value)
    log(user, "INSERT", "inv_dia_term", created.pk, f"{kind}: {value}")
    return created


def rename_term(user, t, value):
    require(user, INV_MASTERS, "Only a role that edits settings can rename a value.")
    value = (value or "").strip()
    if not value:
        raise ServiceError("A value cannot be blank.")
    if DiamondTerm.objects.filter(kind=t.kind, value=value).exclude(pk=t.pk).exists():
        raise ServiceError(f"{value} is already on the list.")
    old, t.value = t.value, value
    t.save(update_fields=["value"])
    for other in DiamondTerm.objects.filter(kind=t.kind).exclude(expands_to=""):
        grades = other.expands_to.split()
        if old in grades:
            other.expands_to = " ".join(value if g == old else g for g in grades)
            other.save(update_fields=["expands_to"])
    log(user, "UPDATE", "inv_dia_term", t.pk, f"{t.kind}: {old} → {value}")


def delete_term(user, t):
    require(user, INV_MASTERS, "Only a role that edits settings can delete a value.")
    if _in_use(t):
        raise ServiceError(f"{t.value} is in use and cannot be deleted.")
    log(user, "DELETE", "inv_dia_term", t.pk, f"{t.kind}: {t.value}")
    t.delete()


def set_expansion(user, t, grades_text):
    require(user, INV_MASTERS, "Only a role that edits settings can change an expansion.")
    grades = (grades_text or "").split()
    if not grades and not t.value.startswith("?"):
        raise ServiceError(f"{t.value} has to expand to at least one grade.")
    known = set(DiamondTerm.objects.filter(kind=t.kind, expands_to="").values_list("value", flat=True))
    unknown = [g for g in grades if g not in known]
    if unknown:
        raise ServiceError(f"Not a single {t.kind} grade: {', '.join(unknown)}.")
    t.expands_to = " ".join(grades)
    t.save(update_fields=["expands_to"])
    log(user, "UPDATE", "inv_dia_term", t.pk, f"{t.value} expands to {t.expands_to or 'nothing'}")


def save_code(user, code, shape, colour, clarity, confirmed, note):
    require(user, INV_MASTERS, "Only a role that edits settings can change an item code.")
    for value, kind in ((shape, DiamondTerm.SHAPE), (colour, DiamondTerm.COLOUR), (clarity, DiamondTerm.CLARITY)):
        if value is not None and value.kind != kind:
            raise ServiceError(f"{value} is not a {kind}.")
    code.shape, code.colour, code.clarity = shape, colour, clarity
    code.confirmed, code.note = bool(confirmed), (note or "").strip()[:120]
    code.save()
    log(user, "UPDATE", "inv_dia_code", code.pk, f"{code.pk}: {shape} / {colour} / {clarity}, confirmed={code.confirmed}")


def _rate(value):
    if value is None:
        return None
    try:
        value = Decimal(value)
    except (InvalidOperation, TypeError, ValueError):
        raise ServiceError("A rate has to be a number from 0 up to ten digits.")
    if not value.is_finite() or value < 0 or abs(value) >= 10 ** 10:
        raise ServiceError("A rate has to be a number from 0 up to ten digits.")
    return value


def set_rate(user, code, size_text, cost_rate, sale_rate, effective_from):
    """A new dated row; nothing overwritten. A row carries both rates, so both rights are needed."""
    require(user, INV_MASTERS, "Only a role that edits settings can set a rate.")
    require(user, VIEW_COST, "A cost you may not see is not yours to set.")
    require(user, VIEW_SALE, "A sale price you may not see is not yours to set.")
    rate = DiamondRate.objects.create(
        code=code, size_text=(size_text or "").strip(), cost_rate=_rate(cost_rate), sale_rate=_rate(sale_rate),
        effective_from=effective_from or timezone.localdate(), set_by=_by(user),
    )
    log(user, "INSERT", "inv_dia_rate", rate.pk, f"{code.pk} {rate.size_text or 'any size'}")
    return rate


def load_ivy_rates(user, fileobj):
    """Sale rates from the IVY export file (not the stock app), per item code and size.

    Only the sale rate: cost comes from the diamond register itself (the owner, 2026-10-01)."""
    require(user, INV_MASTERS, "Only a role that edits settings can load rates.")
    require(user, VIEW_COST, "A cost you may not see is not yours to set.")
    require(user, VIEW_SALE, "A sale price you may not see is not yours to set.")
    try:
        sheet = load_workbook(fileobj, read_only=True, data_only=True).active
    except Exception:
        raise ServiceError("That is not an Excel workbook.")
    codes = set(DiamondCode.objects.values_list("item_code", flat=True))
    found, unknown = {}, set()
    for values in sheet.iter_rows(min_row=IVY_HEADER_ROW + 1, values_only=True):
        values = list(values) + [None] * 40
        code = canonical_code(str(values[IVY_CODE] or ""))
        if not code or values[IVY_SALE] is None:
            continue
        if code not in codes:
            unknown.add(code)
            continue
        size = str(values[IVY_SIZE] or "").strip()
        found[(code, size)] = values[IVY_SALE]
    today = timezone.localdate()
    with transaction.atomic():
        for (code, size), sale in found.items():
            DiamondRate.objects.create(code_id=code, size_text=size, effective_from=today, set_by=_by(user),
                                       sale_rate=_rate(str(sale)))
    log(user, "IMPORT", "inv_dia_rate", "-", f"{len(found)} sale rates from the IVY export, {len(unknown)} unknown codes")
    return {"loaded": len(found), "unknown": len(unknown)}


def save_supplier(user, vendor, code, name, city, terms):
    require(user, INV_MASTERS, "Only a role that edits settings can change suppliers.")
    require(user, VIEW_VENDOR, "A supplier you may not see is not yours to change.")
    code, name = (code or "").strip().upper(), (name or "").strip()
    if not code or not name:
        raise ServiceError("A supplier needs a code and a name.")
    inputs.fits(Vendor, code=code, name=name, city=(city or "").strip(), terms=(terms or "").strip())
    clash = Vendor.objects.filter(code=code)
    if vendor is not None:
        clash = clash.exclude(pk=vendor.pk)
    if clash.exists():
        raise ServiceError(f"Supplier code {code} is already used.")
    vendor = vendor or Vendor()
    vendor.code, vendor.name, vendor.city, vendor.terms = code, name, (city or "").strip(), (terms or "").strip()
    vendor.save()
    log(user, "UPDATE", "vendor", vendor.pk, f"{vendor.code} {vendor.name}")
    return vendor


def set_right(user, role, codename, on):
    """Tick or untick one right for one role. Admins only; the Admin role is never edited."""
    if not (user and user.is_authenticated and user.is_admin()):
        raise PermissionDenied("Only an admin can change user rights.")
    if role == "ADMIN":
        raise ServiceError("The Admin role always holds every right.")
    if role not in ROLE_ORDER or codename not in dict(RIGHTS):
        raise ServiceError("Unknown role or right.")
    group = Group.objects.get(name=role)
    permission = Permission.objects.get(codename=codename, content_type__app_label="accounts")
    (group.permissions.add if on else group.permissions.remove)(permission)
    log(user, "UPDATE", "auth_group", group.pk, f"{role}: {codename} {'on' if on else 'off'}")
