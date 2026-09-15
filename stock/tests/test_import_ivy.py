"""Reading the IVY export.

The rule most likely to break is that the five material bands run down a
block independently — a stone line on row 3 of a block is not related to the
diamond line on row 1. These tests exist mostly to keep that true.
"""
from decimal import Decimal

from stock.importers import ivy
from stock.tests.fixtures_ivy import build_workbook


def test_blocks_start_where_style_no_is_filled():
    pieces = ivy.parse(build_workbook())
    assert [p.jewel_code for p in pieces] == ["24P00088", "24P00111", "24P00095"]
    assert [p.row_no for p in pieces] == [4, 5, 8]


def test_header_fields_come_off_the_parent_row():
    first = ivy.parse(build_workbook())[0]
    assert first.style_code == "ER00502"
    assert first.category == "Earring"
    assert first.collection == "Norna"
    assert first.vendor == "Infinity Venture [IVY]"
    assert first.metal_purity == "18K"
    assert first.diamond_quality == "FGH VS-SI"
    # 'Finish Goods' is normalised to the value the model expects
    assert first.stock_type == "FINISH_GOODS"


def test_misc_remarks_is_read_as_the_fg_date():
    """Column Q is labelled 'Misc Remarks' but the export puts a date in it."""
    first = ivy.parse(build_workbook())[0]
    assert first.fg_date.year == 2025
    assert first.fg_date.month == 11
    assert first.fg_date.day == 7


def test_bands_are_read_independently_down_the_block():
    """The ragged product: 3 diamond, 1 metal, 2 stone lines."""
    ragged = ivy.parse(build_workbook())[1]
    by_band = {}
    for line in ragged.lines:
        by_band.setdefault(line.band, []).append(line)
    assert len(by_band["diamond"]) == 3
    assert len(by_band["metal"]) == 1
    assert len(by_band["stone"]) == 2
    # the stone on block-row 2 must not have been dropped when metal ran out
    assert [s.code for s in by_band["stone"]] == ["SP01C", "PS01W"]
    assert [d.code for d in by_band["diamond"]] == [
        "DRIJ SI-I", "FPL", "DRFGH SI-I",
    ]


def test_line_figures_keep_their_precision():
    ragged = ivy.parse(build_workbook())[1]
    metal = next(l for l in ragged.lines if l.band == "metal")
    assert metal.code == "G14K"
    assert metal.qty == Decimal("8.228")      # net weight, not gross
    assert metal.cost_rate == Decimal("4319")
    stone = next(l for l in ragged.lines if l.code == "PS01W")
    assert stone.pcs == 2
    assert stone.qty == Decimal("44.5")
    assert stone.sale_rate == Decimal("90")


def test_totals_and_source_figures_come_off_the_parent_row():
    ragged = ivy.parse(build_workbook())[1]
    assert ragged.src_cost_price == Decimal("84438")
    assert ragged.src_sale_price == Decimal("227427.18")
    assert ragged.src_net_wt_gm == Decimal("8.228")


def test_charge_lines_are_read_from_the_last_band():
    first = ivy.parse(build_workbook())[0]
    charge = next(l for l in first.lines if l.band == "charge")
    assert charge.code == "EC"
    assert charge.name == "Extra Charges"
    assert charge.sale_amount == Decimal("5400")


def test_a_workbook_with_the_wrong_headers_is_rejected():
    from openpyxl import load_workbook
    stream = build_workbook()
    book = load_workbook(stream)
    book["Sheet1"]["D3"] = "Some Other Column"
    import io
    broken = io.BytesIO()
    book.save(broken)
    broken.seek(0)
    problems = ivy.header_problems(broken)
    assert problems
    assert "JewelCode" in problems[0]


def test_a_correct_workbook_has_no_header_problems():
    assert ivy.header_problems(build_workbook()) == []


def test_the_footer_row_does_not_leak_into_the_last_block():
    """The last block runs to EOF, so the export footer sits inside its range."""
    last = ivy.parse(build_workbook())[-1]
    assert last.jewel_code == "24P00095"
    assert len(last.lines) == 2          # one diamond, one metal — and no footer


# ── guessing new materials ───────────────────────────────────────────────
import pytest

from stock.enums import ChargeBasis, Uom
from stock.importers import guess
from stock.importers.ivy import ParsedLine

pytestmark = pytest.mark.django_db


def _line(band, code, name=""):
    return ParsedLine(band=band, code=code, name=name or code)


def test_a_diamond_code_becomes_a_diamond_in_carats():
    fields, problem = guess.material_fields(_line("diamond", "DRFGH SI-I"))
    assert problem is None
    assert fields["category_id"] == "DIAMOND"
    assert fields["default_uom"] == Uom.CT


def test_a_foil_polki_code_becomes_polki_not_diamond():
    fields, problem = guess.material_fields(_line("diamond", "FPL", "Foil Polki"))
    assert problem is None
    assert fields["category_id"] == "POLKI"


def test_a_gold_code_resolves_its_metal_and_purity(rates):
    fields, problem = guess.material_fields(_line("metal", "G18K", "Gold18K"))
    assert problem is None
    assert fields["category_id"] == "METAL"
    assert fields["metal_id"] == "GOLD"
    assert fields["purity_factor"] == pytest.approx(0.75, abs=0.0001)
    assert fields["default_uom"] == Uom.GM


def test_a_silver_code_resolves_to_silver(rates):
    fields, problem = guess.material_fields(_line("metal", "S925", "Silver925"))
    assert problem is None
    assert fields["metal_id"] == "SILVER"


def test_a_purity_that_does_not_exist_is_a_blocker_not_a_guess(rates):
    """G12K: MetalPurity has no 12K row, and inventing one would be a lie."""
    fields, problem = guess.material_fields(_line("metal", "G12K", "Gold12K"))
    assert problem is not None
    assert "12K" in problem


def test_an_unreadable_metal_code_is_a_blocker(rates):
    """CJ 'Customer Jewelry999' has no G/S prefix to read a metal from."""
    _, problem = guess.material_fields(_line("metal", "CJ", "Customer Jewelry999"))
    assert problem is not None


def test_a_stone_becomes_a_setting_stone():
    fields, problem = guess.material_fields(_line("stone", "SP01C", "Semi Precious"))
    assert problem is None
    assert fields["category_id"] == "SETTING"


def test_an_other_band_line_mints_a_code_from_its_name():
    fields, problem = guess.material_fields(_line("other", "Lakh", "Lakh"))
    assert problem is None
    assert fields["item_code"] == "OTH-LAKH"
    assert fields["category_id"] == "OTHER"


def test_a_charge_becomes_labour():
    fields, problem = guess.material_fields(_line("charge", "EC", "Extra Charges"))
    assert problem is None
    assert fields["category_id"] == "LABOUR"


def test_metal_lines_are_by_qty_never_by_net_metal_weight():
    """BY_NET_METAL_WT would snap every metal line to the piece total."""
    basis, uom = guess.bom_basis_and_uom(_line("metal", "G18K"), "METAL")
    assert basis == ChargeBasis.BY_QTY
    assert uom == Uom.GM


def test_charge_lines_are_flat_so_the_rate_is_the_amount():
    basis, uom = guess.bom_basis_and_uom(_line("charge", "EC"), "LABOUR")
    assert basis == ChargeBasis.FLAT


# ── the handover layout ──────────────────────────────────────────────────
from stock.tests.fixtures_ivy import build_handover


def test_the_handover_layout_is_accepted():
    """Its header block is shifted, but it is still a stock export."""
    assert ivy.header_problems(build_handover()) == []


def test_category_is_read_by_name_not_by_column():
    """The bug this guards: E became 'Image Name', so Category moved to F.

    Read positionally, every piece took the style code as its category and the
    real category as its sub-category — which creates one junk Category per
    style and files nothing where it belongs.
    """
    pieces = ivy.parse(build_handover())
    by_code = {p.jewel_code: p for p in pieces}
    assert by_code["24P00088"].category == "Earring"
    assert by_code["24P00095"].category == "Ring"
    assert "RG" not in by_code["24P00095"].category
    assert by_code["24P00088"].sub_category == "STUDS"


def test_both_layouts_yield_the_same_pieces():
    """The only difference between the two files is where the columns sit."""
    old = {(p.jewel_code, p.category, p.sub_category, len(p.lines))
           for p in ivy.parse(build_workbook())}
    new = {(p.jewel_code, p.category, p.sub_category, len(p.lines))
           for p in ivy.parse(build_handover())}
    assert old == new


def test_a_workbook_missing_a_required_column_is_refused():
    from openpyxl import load_workbook
    import io

    book = load_workbook(build_handover())
    book["Sheet1"]["D3"] = "Something Else"      # JewelCode gone
    broken = io.BytesIO()
    book.save(broken)
    broken.seek(0)
    problems = ivy.header_problems(broken)
    assert problems and "JewelCode" in problems[0]


# ── the same piece listed twice ──────────────────────────────────────────
def _twice(first_date, second_date):
    """One jewel code as two blocks, inwarded on the given dates."""
    from stock.tests.fixtures_ivy import THREE_PRODUCTS

    header, block = THREE_PRODUCTS[0]
    return [
        ({**header, 8: first_date, 2: 1}, block),
        ({**header, 8: second_date, 2: 2}, block),
    ]


def test_the_later_inward_date_wins():
    """A handover file lists a piece again each time it moves."""
    from datetime import datetime

    pieces = ivy.parse(build_workbook(
        products=_twice(datetime(2026, 5, 9), datetime(2026, 7, 25))
    ))
    assert len(pieces) == 1, "two blocks for one jewel code must fold into one"
    assert pieces[0].inw_date.month == 7
    assert pieces[0].superseded == 1


def test_the_later_date_wins_whichever_order_it_appears_in():
    from datetime import datetime

    pieces = ivy.parse(build_workbook(
        products=_twice(datetime(2026, 7, 25), datetime(2026, 5, 9))
    ))
    assert len(pieces) == 1
    assert pieces[0].inw_date.month == 7, "order in the file must not decide it"


def test_a_row_with_no_date_never_beats_one_that_has_a_date():
    from datetime import datetime

    pieces = ivy.parse(build_workbook(products=_twice(datetime(2026, 5, 9), None)))
    assert len(pieces) == 1
    assert pieces[0].inw_date is not None


def test_a_file_with_no_duplicates_folds_nothing():
    pieces = ivy.parse(build_workbook())
    assert len(pieces) == 3
    assert all(p.superseded == 0 for p in pieces)
